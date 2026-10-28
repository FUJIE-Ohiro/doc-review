from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parents[1]/'examples'; OUT.mkdir(exist_ok=True)

def setup(style, eastern='宋体', western='Times New Roman', size=12, bold=False):
    f=style.font; f.name=western; f.size=Pt(size); f.bold=bold
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'),eastern)
    rf.set(qn('w:ascii'),western)
    rf.set(qn('w:hAnsi'),western)

def setfirst(p, chars=200):
    pPr=p._p.get_or_add_pPr()
    ind=pPr.find(qn('w:ind'))
    if ind is None:
        ind=OxmlElement('w:ind'); pPr.append(ind)
    ind.set(qn('w:firstLineChars'),str(chars))

def mk(name,bad=False):
    d=Document()
    s=d.sections[0];s.page_width=Cm(21.0);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2.54)
    styles=d.styles
    normal=styles['Normal']; setup(normal)
    normal.paragraph_format.line_spacing=1
    normal.paragraph_format.space_before=Pt(0); normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for nm,lev in [('Title',0),('Subtitle',0),('Heading 1',1),('Heading 2',2)]:
        sty=styles[nm]
        setup(sty,'黑体' if nm=='Title' else '宋体','Arial' if nm=='Title' else 'Times New Roman',16 if nm=='Title' else 12,nm in ['Title','Heading 1','Heading 2'])
        sty.paragraph_format.line_spacing=1
        sty.paragraph_format.space_before=Pt(0 if nm=='Title' else 12 if lev==1 else 6)
        sty.paragraph_format.space_after=Pt(12 if nm=='Title' else 6 if lev==1 else 0)
        sty.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if nm=='Title' else WD_ALIGN_PARAGRAPH.LEFT
    if bad:
        s.left_margin=Cm(2)
    title=d.add_paragraph('磁性材料领域的三篇论文比较分析', style='Title')
    # Title may cause inherited first line indent? no.
    p=d.add_paragraph('张三');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph('清华大学材料学院');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(12)
    def heading(s,level):
        p=d.add_paragraph(s,style='Heading 1' if level==1 else 'Heading 2')
        return p
    def body(s):
        p=d.add_paragraph(s)
        setfirst(p)
        return p
    heading('1. 研究领域与论文选择',1)
    body('本项目面向磁性材料的实验设计，选择三个发表水平不同、技术路线具有共同点的研究作为分析对象。')
    heading('2. 高水平论文分析',1)
    heading('2.1 研究问题与科学贡献',2)
    body('该论文通过系统研究突出其创新性，同时保留实验的关键参数，以评价结论的可重复性。')
    heading('3. 中等水平论文分析',1)
    heading('3.1 研究问题与科学贡献',2)
    bb=body('通过对比发现论文中应用的方法十分成熟，但缺少对所有局限性的深入讨论。')
    if bad:
        bb.paragraph_format.line_spacing=1.5
        bb.runs[0].font.size=Pt(10.5)
    heading('4. 相对普通水平论文分析',1)
    body('论文包含常规材料性能表征，但缺少实验误差及现有研究的严谨对比。')
    heading('5. 比较分析与总结',1)
    heading('5.1 三篇论文的横向比较',2)
    body('通过科研深度、英文写作方面进行比较，论文A的逻辑组织明显优于论文B和论文C。')
    heading('5.2 总结与启示',2)
    body('比较结果说明科技论文不仅要求结果可靠，也要强调证据和清晰的论证结构。')
    heading('6. 参考文献',1)
    rp=d.add_paragraph('[1] Smith J. Magnetic materials. Journal of Science, 2025. https://doi.org/10.0000/example1')
    rp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    rp=d.add_paragraph('[2] Huang K. Film interfaces. Journal of Magnetism, 2024. https://doi.org/10.0000/example2')
    rp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    # For reference matching, put cites in text; re-save included in main body
    d.paragraphs[4].add_run(' 本文分析依据见[1]与[2]。')
    d.save(OUT/name)
mk('01_符合规范的测试稿.docx')
mk('02_故意设置错误的测试稿.docx',True)
print('fixtures',list(OUT.glob('*.docx')))